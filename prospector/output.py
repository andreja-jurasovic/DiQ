"""Write results to the Google Sheet (one tab per zone, upsert by key) and to CSV.

Columns follow "Kategorije informacija 210926.xlsx", with Veličina split into Promet
and Broj zaposlenih. Status is the salesperson's: once filled in, a re-run never
overwrites it (except for companies listed in customers.csv)."""
import csv
import json
import re
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import gspread
from gspread.utils import rowcol_to_a1

from .config import OUTPUT_DIR, categories, env, scoring

COLUMNS = [
    "Županija", "Firma", "Adresa", "Status", "Segment", "Promet", "Broj zaposlenih", "Web stranica",
    "Kontakt ime", "Kontakt mail", "Kontakt broj tel. ", "Score", "Prio", "Kategorije",
    "Ključne riječi", "Brendovi na webu", "eMajstor profil", "eMajstor usluge", "Izvori", "OIB",
    "Obrtni registar", "Status obrta", "MBO obrta", "Djelatnost obrta", "Izvor prometa",
    "Bilješke", "Datum obrade", "Ažurirano", "Ključ",
]
MANUAL = {"Status", "Bilješke"}


def _revenue(p: dict) -> str:
    if p.get("revenue") is not None:
        rev = p["revenue"]
        text = f"{rev / 1e6:.1f}M EUR" if rev >= 1e6 else f"{rev / 1e3:.0f}k EUR"
        if p.get("revenue_year"):
            text += f" ({p['revenue_year']}.)"
        if p.get("revenue_source") == "procjena (web)":
            text = "~" + text
        return text
    if not p.get("oib"):
        return "Nema javnog financijskog podatka"
    return "Nema podatka"


def _employees(p: dict) -> str:
    return str(int(p["employees"])) if p.get("employees") is not None else ""


def _error_note(error: str | None) -> str:
    if not error:
        return ""
    reason = "nema Anthropic kredita" if "credit balance" in error else error.split(":")[0]
    return f"NIJE KVALIFICIRANO ({reason}) - pokreni ponovno"


def _contact_names(p: dict) -> str:
    names = [f"{c['name']} ({c['role']})" if c.get("role") else c["name"] for c in p.get("contacts", [])]
    contact_name = p.get("contact_name")
    if contact_name and not any(contact_name.casefold() in n.casefold() for n in names):
        names.insert(0, contact_name)
    return " / ".join(names[:3])


def row(p: dict) -> dict:
    tiers = scoring()["tiers"]
    cat_labels = categories()
    return {
        "Županija": (p.get("county") or "").replace(" županija", ""),
        "Firma": p.get("legal_name") or p["name"],
        "Adresa": p.get("register_address") or p.get("address") or p.get("town") or "",
        "Status": p.get("status", "New"),
        # a failed Claude call leaves no segment; say so where the segment would be
        "Segment": p.get("segment") or _error_note(p.get("error")),
        "Promet": _revenue(p),
        "Broj zaposlenih": _employees(p),
        "Web stranica": p.get("website") or "Nema web stranicu",
        "Kontakt ime": _contact_names(p),
        "Kontakt mail": " / ".join(p.get("emails", [])[:3]),
        "Kontakt broj tel. ": p.get("phone_display", ""),
        "Score": "isključeno" if p.get("score") is None else p["score"],
        "Prio": f"{p['tier']}. {tiers[p['tier']]}",
        "Kategorije": ", ".join(cat_labels[c]["label"] for c in p.get("ai_categories", []) if c in cat_labels),
        "Ključne riječi": ", ".join(dict.fromkeys(p.get("category_keywords", []) + p.get("general_keywords", []))),
        "Brendovi na webu": ", ".join(p.get("brands", [])),
        "eMajstor profil": p.get("profile_url") or "",
        "eMajstor usluge": p.get("emajstor_services") or "",
        "Izvori": ", ".join(sorted(p.get("sources", []))),
        "OIB": p.get("oib") or "",
        "Obrtni registar": p.get("obrtni_registar", ""),
        "Status obrta": p.get("obrt_registry_status", ""),
        "MBO obrta": p.get("obrt_registry_mbo", ""),
        "Djelatnost obrta": p.get("obrt_registry_activity", ""),
        "Izvor prometa": " ".join(filter(None, [p.get("revenue_source"), p.get("revenue_url")])),
        "Bilješke": p.get("notes", ""),
        "Datum obrade": p.get("processed_at") or datetime.now().astimezone().isoformat(timespec="minutes"),
        "Ažurirano": date.today().isoformat(),
        "Ključ": f"oib:{p['oib']}" if p.get("oib") else f"name:{(p.get('legal_name') or p['name']).lower()}|{(p.get('town') or '').lower()}",
    }


def write_csv(prospects: list[dict], name: str) -> str:
    path = OUTPUT_DIR / f"{name}.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        for p in prospects:
            w.writerow(row(p))
    return str(path)


def write_xlsx(prospects: list[dict], name: str) -> str:
    """Write a formatted Excel copy of the same full prospect output as the CSV."""
    path = OUTPUT_DIR / f"{name}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospekti"
    ws.append(COLUMNS)
    for prospect in prospects:
        values = row(prospect)
        ws.append([values.get(column, "") for column in COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174A6E")
    for column_index, column in enumerate(COLUMNS, 1):
        sample_values = [str(ws.cell(row_index, column_index).value or "") for row_index in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(max((len(value) for value in sample_values), default=10) + 2, 12), 48)
        ws.column_dimensions[get_column_letter(column_index)].width = width
    wb.save(path)
    return str(path)


def write_json(prospects: list[dict], name: str) -> str:
    """Full run data, so the CSV/Sheet can be rebuilt later without paying for Claude again."""
    path = OUTPUT_DIR / f"{name}.json"
    path.write_text(json.dumps(prospects, ensure_ascii=False, indent=1, default=sorted), encoding="utf-8")
    return str(path)


def _migrate(ws, values: list[list[str]]) -> list[list[str]]:
    """Rewrite an agent tab written with an older column layout into the current one.
    Columns that no longer exist are dropped; new ones start empty."""
    old_header = values[0]
    rows = [dict(zip(old_header, r + [""] * len(old_header))) for r in values[1:]]
    for r in rows:  # old combined column "2.9M EUR (2025.), 22 zaposlenih" -> Promet + Broj zaposlenih
        size = r.get("Veličina (promet, broj radn.)")
        if size and not r.get("Promet"):
            m = re.search(r",\s*(\d+) zaposlenih", size)
            r["Promet"] = size[: m.start()] if m else size
            r["Broj zaposlenih"] = m.group(1) if m else ""
        if not r.get("Segment") and r.get("Call brief", "").startswith(("BadRequestError", "NIJE KVALIFICIRANO")):
            r["Segment"] = _error_note(r["Call brief"])
    new_values = [COLUMNS] + [[r.get(c, "") for c in COLUMNS] for r in rows]
    ws.clear()
    ws.resize(rows=max(1000, len(new_values) + 100), cols=len(COLUMNS))
    ws.update(new_values, "A1", value_input_option="USER_ENTERED")
    return new_values


def write_sheet(prospects: list[dict], tab: str, customer_keys: set[str]) -> str:
    gc = gspread.service_account(filename=env("GOOGLE_SERVICE_ACCOUNT_FILE"))
    try:
        sh = gc.open_by_key(env("GOOGLE_SHEET_ID"))
    except PermissionError:
        raise RuntimeError(
            f"No access to the Google Sheet. Share it (Editor) with {gc.http_client.auth.service_account_email}"
        ) from None
    try:
        ws = sh.worksheet(tab)
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(tab, rows=1000, cols=len(COLUMNS))
    values = ws.get_all_values()
    if values and values[0] and values[0][: len(COLUMNS)] != COLUMNS:
        if "Ključ" in values[0]:
            values = _migrate(ws, values)  # our own tab, older layout
        else:
            # someone else's layout: leave it alone and use a separate tab
            ws = sh.add_worksheet(f"{tab} (agent)", rows=1000, cols=len(COLUMNS))
            values = []
    if not values or not any(values[0]):
        ws.update([COLUMNS], "A1")
        values = [COLUMNS]
    ws.format(f"A1:{rowcol_to_a1(1, len(COLUMNS))}", {"textFormat": {"bold": True}})
    ws.freeze(rows=1)

    key_col = COLUMNS.index("Ključ")
    existing = {r[key_col]: i for i, r in enumerate(values[1:], start=2) if len(r) > key_col and r[key_col]}
    updates, appends = [], []
    for p in prospects:
        new = row(p)
        line = existing.get(new["Ključ"])
        if line is None:
            appends.append([new[c] for c in COLUMNS])
            continue
        old = dict(zip(COLUMNS, values[line - 1] + [""] * len(COLUMNS)))
        for col in MANUAL:
            is_customer = col == "Status" and new["Ključ"] in customer_keys
            if old.get(col) and not is_customer:
                new[col] = old[col]
        updates.append({"range": f"A{line}:{rowcol_to_a1(line, len(COLUMNS))}", "values": [[new[c] for c in COLUMNS]]})
    if updates:
        ws.batch_update(updates, value_input_option="USER_ENTERED")
    if appends:
        ws.append_rows(appends, value_input_option="USER_ENTERED")
    return f"https://docs.google.com/spreadsheets/d/{sh.id}/edit#gid={ws.id} ({len(appends)} new, {len(updates)} updated)"
